from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from app.main import _attach_template_visual_preview, download_template, search_templates
from app.services.export_service import PreviewResult


class _GiteeManager:
    def __init__(self) -> None:
        self.calls: list[tuple[str, str, dict]] = []

    async def start_server(self, name: str) -> None:
        return None

    async def health(self, *, probe: bool = False) -> dict:
        return {"gitee": {"name": "gitee", "status": "ready"}}

    async def call_tool(self, server: str, tool: str, arguments: dict) -> dict:
        self.calls.append((server, tool, dict(arguments)))
        if tool == "gitee_list_resume_templates":
            return {
                "ok": True,
                "data": {
                    "items": [
                        {
                            "template_id": "gitee-0123456789abcdef01234567",
                            "name": "中文技术简历",
                            "file_type": "docx",
                            "language": "zh",
                            "license": "upstream-repository",
                            "license_allowed": True,
                            "repository": "hicv-cn/hicv-word-resume-templates-jianyue-single",
                            "file_path": "templates/02_简约简历/单页/中文技术简历.docx",
                            "source_url": "https://gitee.com/hicv-cn/hicv-word-resume-templates-jianyue-single/blob/main/templates/02_%E7%AE%80%E7%BA%A6%E7%AE%80%E5%8E%86/%E5%8D%95%E9%A1%B5/%E4%B8%AD%E6%96%87%E6%8A%80%E6%9C%AF%E7%AE%80%E5%8E%86.docx",
                        }
                    ]
                },
            }
        return {
            "ok": True,
            "data": {
                "status": "cached",
                "template_id": "gitee-0123456789abcdef01234567",
                "path": "data/templates/cache/template/resume.docx",
                "metadata": {
                    "template_id": "gitee-0123456789abcdef01234567",
                    "repository": "hicv-cn/hicv-word-resume-templates-jianyue-single",
                    "source_url": "https://gitee.com/hicv-cn/hicv-word-resume-templates-jianyue-single",
                    "file_type": "docx",
                    "language": "zh",
                    "license": "upstream-repository",
                    "license_allowed": True,
                },
            },
        }


class _TemplateMirror:
    def __init__(self) -> None:
        self.metadata: list[dict] = []

    def persist_template(self, metadata: dict) -> str:
        self.metadata.append(metadata)
        return str(metadata["template_id"])


class TemplateMcpApiTests(unittest.IsolatedAsyncioTestCase):
    async def test_cached_markdown_preview_is_renderable_not_a_text_excerpt(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            data_root = Path(directory)
            template = data_root / "templates" / "cache" / "github-test" / "resume.md"
            template.parent.mkdir(parents=True)
            template.write_text("# 中文简历\n\n## 项目经历\n", encoding="utf-8")
            app_state = SimpleNamespace(settings=SimpleNamespace(data_root=data_root))
            with patch("app.main.state", return_value=app_state):
                rendered = await _attach_template_visual_preview(
                    {"path": str(template), "file_type": "markdown"},
                    requested=True,
                )

        self.assertEqual(rendered["visual_preview"]["status"], "ready")
        self.assertEqual(rendered["visual_preview"]["mime_type"], "text/markdown")
        self.assertIn("## 项目经历", rendered["visual_preview"]["content"])

    async def test_docx_preview_falls_back_to_safe_browser_page_without_office(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            data_root = Path(directory)
            template = data_root / "templates" / "cache" / "github-test" / "中文简历模板.docx"
            template.parent.mkdir(parents=True)
            document = Document()
            document.add_heading("中文技术简历", level=1)
            document.add_paragraph("个人信息 项目经历 专业技能 教育经历")
            document.save(template)
            app_state = SimpleNamespace(settings=SimpleNamespace(data_root=data_root))
            unavailable = PreviewResult(
                "pdf",
                "dependency_missing",
                None,
                None,
                None,
                ("no office",),
            )
            with (
                patch("app.main.state", return_value=app_state),
                patch("app.main.convert_docx_to_pdf", return_value=unavailable),
            ):
                rendered = await _attach_template_visual_preview(
                    {"path": str(template), "file_type": "docx"},
                    requested=True,
                )

        visual = rendered["visual_preview"]
        self.assertEqual(visual["status"], "ready")
        self.assertEqual(visual["mime_type"], "text/html")
        self.assertTrue(visual["approximate_layout"])
        self.assertIn("docx-page", visual["content"])

    async def test_search_and_download_use_gitee_mcp(self) -> None:
        manager = _GiteeManager()
        mirror = _TemplateMirror()
        app_state = SimpleNamespace(
            mcp_manager=manager,
            mcp_startup_errors={},
            db_mirror=mirror,
        )
        with (
            patch("app.main.state", return_value=app_state),
            patch("app.main._gitee_token", return_value="test-token"),
        ):
            searched = await search_templates({"source": "gitee"}, None)
            self.assertEqual(searched["source"], "gitee_mcp")
            item = searched["items"][0]
            downloaded = await download_template(
                item["template_id"],
                {"file_path": item["file_path"]},
                None,
            )

        self.assertEqual(downloaded["source"], "gitee_mcp")
        self.assertEqual(mirror.metadata[0]["license"], "upstream-repository")
        self.assertEqual(
            [tool for _, tool, _ in manager.calls],
            ["gitee_list_resume_templates", "gitee_download_template"],
        )


if __name__ == "__main__":
    unittest.main()
