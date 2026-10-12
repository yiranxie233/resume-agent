from __future__ import annotations

import base64
import io
import tempfile
import unittest
import urllib.error
from pathlib import Path
from unittest.mock import patch

from docx import Document

from app.services.gitee_template_service import (
    GITEE_BRANCH,
    GITEE_REPOSITORY,
    GITEE_TEMPLATE_PATH,
    GiteeAPIError,
    _request_json,
    cache_gitee_template,
    list_gitee_templates,
    template_id_for_gitee,
)


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("中文技术简历", level=1)
    document.add_paragraph(
        "个人信息 联系方式 求职目标 教育经历 工作经历 项目经历 专业技能 "
        "熟悉 Python FastAPI LangGraph PostgreSQL Docker，具备完整项目交付经验。"
    )
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class GiteeTemplateServiceTests(unittest.TestCase):
    def test_listing_is_locked_to_direct_docx_children(self) -> None:
        directory = GITEE_TEMPLATE_PATH
        payload = [
            {
                "type": "file",
                "name": "01.docx",
                "path": f"{directory}/01.docx",
                "sha": "abc",
                "size": 1200,
            },
            {
                "type": "file",
                "name": "readme.md",
                "path": f"{directory}/readme.md",
                "sha": "def",
                "size": 10,
            },
            {
                "type": "dir",
                "name": "nested",
                "path": f"{directory}/nested",
            },
            {
                "type": "file",
                "name": "outside.docx",
                "path": "templates/outside.docx",
                "sha": "ghi",
                "size": 100,
            },
        ]
        with patch(
            "app.services.gitee_template_service._request_json",
            return_value=payload,
        ):
            result = list_gitee_templates(token="private-token", page=1, per_page=5)

        self.assertEqual(result["repository"], GITEE_REPOSITORY)
        self.assertEqual(result["branch"], GITEE_BRANCH)
        self.assertEqual(result["template_path"], directory)
        self.assertEqual(len(result["items"]), 1)
        self.assertEqual(result["items"][0]["file_path"], f"{directory}/01.docx")
        self.assertTrue(result["items"][0]["template_id"].startswith("gitee-"))

    def test_cache_accepts_numeric_docx_name_from_trusted_directory(self) -> None:
        path = f"{GITEE_TEMPLATE_PATH}/001.docx"
        data = _docx_bytes()
        payload = {
            "type": "file",
            "path": path,
            "sha": "remote-sha",
            "content": base64.b64encode(data).decode("ascii"),
        }
        with tempfile.TemporaryDirectory() as directory, patch(
            "app.services.gitee_template_service._request_json",
            return_value=payload,
        ):
            result = cache_gitee_template(
                file_path=path,
                cache_root=Path(directory),
                token="private-token",
            )
            metadata_text = (Path(directory) / result["template_id"] / "metadata.json").read_text(
                encoding="utf-8"
            )

        self.assertEqual(result["template_id"], template_id_for_gitee(path))
        self.assertEqual(result["metadata"]["file_type"], "docx")
        self.assertNotIn("private-token", metadata_text)
        self.assertNotIn("access_token", metadata_text)

    def test_http_error_redacts_query_token(self) -> None:
        secret = "gitee-secret-value"
        url = f"https://gitee.com/api/v5/user?access_token={secret}"
        error = urllib.error.HTTPError(url, 401, "unauthorised", {}, None)
        with patch("urllib.request.urlopen", side_effect=error):
            with self.assertRaises(GiteeAPIError) as caught:
                _request_json(url)
        self.assertEqual(caught.exception.status_code, 401)
        self.assertNotIn(secret, str(caught.exception))

    def test_arbitrary_repository_path_is_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "固定 Gitee 目录"):
            template_id_for_gitee("templates/other/evil.docx")


if __name__ == "__main__":
    unittest.main()
