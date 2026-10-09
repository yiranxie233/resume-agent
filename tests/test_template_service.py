from __future__ import annotations

import hashlib
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from app.mcp import github_server
from app.services import template_service


CHINESE_TEMPLATE = (
    "# 中文技术简历\n\n## 个人信息\n姓名：[待填写]\n\n"
    "## 教育经历\n[待填写]\n\n## 专业技能\n[待填写]\n\n"
    "## 项目经历\n[待填写]\n"
).encode("utf-8")

NOT_A_TEMPLATE = (
    "# 每周简讯\n\n本文讨论如何编写个人信息、项目经历和技能清单，"
    "但它是一篇说明文章而不是可以填写的简历。\n"
).encode("utf-8")


def git_blob_sha(data: bytes) -> str:
    return hashlib.sha1(f"blob {len(data)}\0".encode("ascii") + data).hexdigest()


class TemplateServiceTests(unittest.TestCase):
    def test_docx_browser_preview_is_paginated_and_escapes_content(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "中文简历模板.docx"
            document = Document()
            document.add_heading("张三 <script>", level=1)
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "项目经历"
            table.cell(0, 1).text = "Python 平台"
            document.save(path)
            rendered = template_service.render_docx_html_preview(path.read_bytes())

        self.assertIn("docx-page", rendered)
        self.assertIn("<table>", rendered)
        self.assertIn("&lt;script&gt;", rendered)
        self.assertNotIn("张三 <script>", rendered)

    def test_template_download_retries_transient_raw_github_timeout(self) -> None:
        source = "https://raw.githubusercontent.com/acme/cv/main/resume.md"
        with patch.object(
            template_service,
            "_request_bytes",
            side_effect=[
                TimeoutError("slow route"),
                (CHINESE_TEMPLATE, source, "text/markdown"),
            ],
        ) as request_bytes, patch.object(template_service.time, "sleep"):
            data, final_url, content_type = template_service._request_template_bytes(
                source
            )

        self.assertEqual(data, CHINESE_TEMPLATE)
        self.assertEqual(final_url, source)
        self.assertEqual(content_type, "text/markdown")
        self.assertEqual(request_bytes.call_count, 2)

    def test_article_with_resume_keywords_is_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "不是简历模板|内容校验"):
            template_service._preview_metadata(NOT_A_TEMPLATE, "docs/weekly.md")

    def test_license_is_fetched_and_allowlisted(self) -> None:
        payload = {
            "license": {"spdx_id": "MIT", "name": "MIT License"},
            "html_url": "https://github.com/acme/cv/blob/main/LICENSE",
        }
        with patch.object(template_service, "_request_json", return_value=payload) as request:
            result = template_service.get_github_repository_license(
                "acme/cv", token="session-secret"
            )

        self.assertTrue(result["license_allowed"])
        self.assertTrue(result["allows_modification"])
        self.assertNotIn("session-secret", json.dumps(result))
        self.assertEqual(request.call_args.kwargs["token"], "session-secret")

    def test_attribution_license_is_not_allowed(self) -> None:
        payload = {
            "license": {"spdx_id": "CC-BY-4.0", "name": "Creative Commons Attribution"}
        }
        with patch.object(template_service, "_request_json", return_value=payload):
            result = template_service.get_github_repository_license("acme/cv")

        self.assertFalse(result["license_allowed"])
        self.assertEqual(result["status"], "blocked")

    def test_search_returns_verified_file_level_chinese_candidate(self) -> None:
        sha = git_blob_sha(CHINESE_TEMPLATE)

        def json_response(url: str, **_: object) -> object:
            if "/search/repositories?" in url:
                return {
                    "total_count": 1,
                    "items": [
                        {
                            "full_name": "acme/chinese-cv",
                            "html_url": "https://github.com/acme/chinese-cv",
                            "stargazers_count": 321,
                            "updated_at": "2026-09-01T00:00:00Z",
                            "default_branch": "main",
                            "license": {"spdx_id": "MIT"},
                        }
                    ],
                }
            if url.endswith("/license"):
                return {"license": {"spdx_id": "MIT", "name": "MIT License"}}
            if "/git/trees/" in url:
                return {
                    "tree": [
                        {
                            "type": "blob",
                            "path": "templates/中文简历.md",
                            "sha": sha,
                            "size": len(CHINESE_TEMPLATE),
                        }
                    ]
                }
            raise AssertionError(url)

        with (
            patch.object(template_service, "_request_json", side_effect=json_response),
            patch.object(
                template_service,
                "_request_bytes",
                return_value=(
                    CHINESE_TEMPLATE,
                    "https://raw.githubusercontent.com/acme/chinese-cv/main/"
                    "templates/%E4%B8%AD%E6%96%87%E7%AE%80%E5%8E%86.md",
                    "text/plain",
                ),
            ),
        ):
            result = template_service.search_github_templates(per_page=5)

        self.assertEqual(len(result["items"]), 1)
        item = result["items"][0]
        self.assertEqual(item["repository"], "acme/chinese-cv")
        self.assertEqual(item["file_type"], "markdown")
        self.assertEqual(item["language"], "zh")
        self.assertTrue(item["license_allowed"])
        self.assertTrue(item["preview_available"])
        self.assertEqual(item["remote_version"], f"git:{sha}")
        self.assertTrue(item["source_url"].startswith("https://raw.githubusercontent.com/"))

    def test_docx_search_defers_large_file_validation_to_download_stage(self) -> None:
        def json_response(url: str, **_: object) -> object:
            if "/search/repositories?" in url:
                return {
                    "total_count": 1,
                    "items": [
                        {
                            "full_name": "acme/chinese-resume-templates",
                            "html_url": "https://github.com/acme/chinese-resume-templates",
                            "stargazers_count": 999,
                            "default_branch": "main",
                            "license": {"spdx_id": "MIT"},
                        }
                    ],
                }
            if url.endswith("/license"):
                return {"license": {"spdx_id": "MIT", "name": "MIT License"}}
            if "/git/trees/" in url:
                return {
                    "tree": [
                        {
                            "type": "blob",
                            "path": "中文/软件开发工程师简历模板.docx",
                            "sha": "abc123",
                            "size": 350_000,
                        }
                    ]
                }
            raise AssertionError(url)

        with (
            patch.object(template_service, "_request_json", side_effect=json_response),
            patch.object(template_service, "_request_bytes") as request_bytes,
        ):
            result = template_service.search_github_templates(per_page=1)

        self.assertEqual(len(result["items"]), 1)
        self.assertEqual(
            result["items"][0]["preview_metadata"]["validation_stage"], "download"
        )
        request_bytes.assert_not_called()

    def test_download_rejects_non_github_url_before_network_access(self) -> None:
        with tempfile.TemporaryDirectory() as temporary, patch.object(
            template_service.urllib.request, "urlopen"
        ) as request:
            with self.assertRaisesRegex(ValueError, "GitHub"):
                template_service.download_github_file(
                    "https://example.com/template.md", root=Path(temporary)
                )
        request.assert_not_called()

    def test_cache_metadata_excludes_token_and_update_check_does_not_overwrite(self) -> None:
        source_url = "https://raw.githubusercontent.com/acme/cv/main/resume.md"
        initial_sha = git_blob_sha(CHINESE_TEMPLATE)
        license_result = {
            "repository": "acme/cv",
            "license": "MIT",
            "license_allowed": True,
            "allows_local_use": True,
            "allows_modification": True,
            "requires_output_attribution": False,
            "policy_version": "github-template-license-v1",
            "license_name": "MIT License",
            "status": "allowed",
            "license_url": None,
        }
        with (
            tempfile.TemporaryDirectory() as temporary,
            patch.object(
                template_service, "get_github_repository_license", return_value=license_result
            ),
            patch.object(
                template_service,
                "_request_bytes",
                return_value=(CHINESE_TEMPLATE, source_url, "text/plain"),
            ),
        ):
            root = Path(temporary)
            cached = template_service.cache_github_template(
                source_url=source_url,
                cache_root=root,
                token="never-write-this-token",
                remote_version=f"git:{initial_sha}",
            )
            metadata_path = root / cached["template_id"] / "metadata.json"
            metadata_text = metadata_path.read_text(encoding="utf-8")
            cached_bytes = Path(cached["path"]).read_bytes()
            self.assertNotIn("never-write-this-token", metadata_text)
            self.assertNotIn("token", metadata_text.lower())

            with (
                patch.object(
                    template_service,
                    "get_github_repository_license",
                    return_value=license_result,
                ),
                patch.object(
                    template_service,
                    "get_github_file_version",
                    return_value={
                        "repository": "acme/cv",
                        "file_path": "resume.md",
                        "ref": "main",
                        "remote_version": "git:new-version",
                        "remote_size_bytes": 99,
                        "source_url": source_url,
                        "url": "https://github.com/acme/cv/blob/main/resume.md",
                    },
                ),
            ):
                update = template_service.check_github_template_update(
                    cache_root=root, template_id=cached["template_id"], token="ephemeral"
                )

            self.assertEqual(update["status"], "update_available")
            self.assertFalse(update["auto_overwrite"])
            self.assertEqual(Path(cached["path"]).read_bytes(), cached_bytes)

    def test_cached_preview_cannot_read_outside_cache_root(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary) / "cache"
            root.mkdir()
            outside = Path(temporary) / "resume.md"
            outside.write_bytes(CHINESE_TEMPLATE)
            with self.assertRaisesRegex(ValueError, "缓存目录"):
                template_service.preview_cached_template(outside, cache_root=root)


class GithubMCPTests(unittest.TestCase):
    def test_search_token_only_comes_from_current_call(self) -> None:
        captured: dict[str, object] = {}

        def fake_search(**kwargs: object) -> dict[str, object]:
            captured.update(kwargs)
            return {"items": [], "page": 1, "has_more": False}

        with (
            patch.dict(os.environ, {"GITHUB_TOKEN": "must-not-be-used"}),
            patch.object(github_server, "search_github_templates", side_effect=fake_search),
        ):
            result = github_server.search({"query": "中文模板"})

        self.assertTrue(result["ok"])
        self.assertIsNone(captured["token"])
        self.assertNotIn("must-not-be-used", json.dumps(result, ensure_ascii=False))

    def test_license_tool_rejects_unknown_license(self) -> None:
        value = {
            "repository": "acme/cv",
            "license": None,
            "license_allowed": False,
            "status": "unknown",
        }
        with patch.object(github_server, "get_github_repository_license", return_value=value):
            result = github_server.license_info({"repository": "acme/cv"})
        self.assertFalse(result["ok"])
        self.assertEqual(result["error_code"], "template_license_unknown")


if __name__ == "__main__":
    unittest.main()
