from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from app.core.schemas import CandidateProject
from app.services.export_service import (
    apply_markdown_template,
    compress_markdown,
    estimate_pages,
    export_docx,
    export_text_pdf,
    render_markdown,
)
from app.ui.streamlit_app import _candidate_plain_text, _preview_matches_task, _selected_candidate


def _candidate(candidate_id: str, title: str) -> CandidateProject:
    return CandidateProject(
        candidate_id=candidate_id,
        candidate_slot_id=f"slot-{candidate_id[-1]}",
        title=title,
        introduction=f"{title}简介",
        tech_stack=["Python"],
        solutions=[f"{title}方案"],
        results=[f"{title}结果"],
    )


class ExportServiceTests(unittest.TestCase):
    def test_render_markdown_uses_explicit_selected_candidate(self) -> None:
        first = _candidate("candidate-1", "首个候选")
        second = _candidate("candidate-2", "用户选择候选")

        rendered = render_markdown(
            job_title="Python 工程师",
            candidates=[first, second],
            selected_candidate_id=second.candidate_id,
        )

        self.assertIn("用户选择候选", rendered)
        self.assertNotIn("首个候选", rendered)

    def test_compression_changes_layout_before_rewriting(self) -> None:
        content = "正文" * 1100
        self.assertEqual(estimate_pages(content), 2)

        compressed = compress_markdown(content, allow_rewrite=False)

        self.assertEqual(compressed.content, content)
        self.assertFalse(compressed.rewritten)
        self.assertEqual(compressed.page_count, 1)
        self.assertGreater(compressed.density, 1.0)

    def test_aggressive_compression_keeps_every_character_and_adds_compact_reflow(self) -> None:
        content = "正文" * 2000

        standard = compress_markdown(content, allow_rewrite=False, level="standard")
        aggressive = compress_markdown(content, allow_rewrite=False, level="aggressive")

        self.assertEqual(aggressive.content, content)
        self.assertFalse(aggressive.rewritten)
        self.assertTrue(aggressive.compact_reflow)
        self.assertEqual(aggressive.level, "aggressive")
        self.assertLess(aggressive.page_count, standard.page_count)

    def test_template_output_contains_base_facts_and_selected_project(self) -> None:
        selected = _candidate("candidate-2", "智能体岗位匹配平台")
        template = "# {{姓名}}\n\n## 个人信息\n\n[待补充]\n\n## 项目经历\n\n[待补充]\n"

        rendered = render_markdown(
            job_title="AI 工程师",
            candidates=[selected],
            selected_candidate_id=selected.candidate_id,
            base_facts={
                "name": "张三",
                "phone": "13800000000",
                "email": "zhangsan@example.com",
                "objective": "AI 工程师",
            },
            template_content=template,
        )

        self.assertIn("# 张三", rendered)
        self.assertIn("13800000000", rendered)
        self.assertIn("zhangsan@example.com", rendered)
        self.assertIn("智能体岗位匹配平台", rendered)
        self.assertNotIn("## 项目经历\n\n[待补充]", rendered)

    def test_markdown_template_appends_project_when_section_is_absent(self) -> None:
        rendered = apply_markdown_template(
            "# {{name}}\n",
            base_facts={"name": "李四"},
            project_content="### 新项目\n\n项目内容",
            job_title="Python 工程师",
        )

        self.assertIn("# 李四", rendered)
        self.assertIn("## 项目经历", rendered)
        self.assertIn("### 新项目", rendered)

    def test_plain_text_candidate_is_complete_and_copy_ready(self) -> None:
        text = _candidate_plain_text(
            _candidate("candidate-3", "可复制项目").model_dump(mode="json")
        )

        for label in ("项目简介：", "技术栈：", "解决方案：", "结果："):
            self.assertIn(label, text)
        self.assertNotIn("```", text)

    def test_docx_template_keeps_layout_and_replaces_supported_placeholders(self) -> None:
        from docx import Document
        from docx.shared import Cm

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "template.docx"
            output = root / "resume.docx"
            document = Document()
            document.sections[0].left_margin = Cm(2.5)
            document.add_paragraph("{{name}}")
            document.add_paragraph("{{resume_content}}")
            document.save(source)

            result = export_docx(
                "# 目标简历\n\n## 项目经历\n\n项目正文",
                target=output,
                template_path=source,
                template_replacements={"name": "张三"},
            )
            rendered = Document(output)

        self.assertEqual(result.status, "preview_pending")
        self.assertAlmostEqual(rendered.sections[0].left_margin.cm, 2.5, places=1)
        text = "\n".join(item.text for item in rendered.paragraphs)
        self.assertIn("张三", text)
        self.assertIn("项目正文", text)
        self.assertNotIn("生成内容", text)

    def test_docx_template_without_english_heading_styles_does_not_crash(self) -> None:
        from docx import Document
        from docx.oxml.ns import qn

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "chinese-template.docx"
            output = root / "resume.docx"
            document = Document()
            for style in list(document.styles.element):
                name = style.find(qn("w:name"))
                if name is not None and name.get(qn("w:val")) in {
                    "heading 1",
                    "heading 2",
                    "heading 3",
                }:
                    document.styles.element.remove(style)
            document.add_paragraph("中文模板正文")
            document.save(source)

            result = export_docx(
                "# 生成简历\n\n## 项目经历\n\n项目正文",
                target=output,
                template_path=source,
            )
            rendered_text = "\n".join(
                paragraph.text for paragraph in Document(output).paragraphs
            )

        self.assertEqual(result.status, "preview_pending")
        self.assertIn("项目正文", rendered_text)
        self.assertNotIn("中文模板正文", rendered_text)

    def test_text_pdf_fallback_preserves_chinese_content(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "resume.pdf"
            result = export_text_pdf(
                "# 中文简历\n\n## 项目经历\n\n智能体项目正文",
                target=target,
            )

            self.assertEqual(result.status, "ready")
            self.assertTrue(target.read_bytes().startswith(b"%PDF"))
            self.assertGreaterEqual(result.page_count or 0, 1)

    def test_ui_resolves_only_the_explicit_selection(self) -> None:
        task = {
            "task_id": "task-1",
            "resume_snapshot_id": "snapshot-1",
            "selected_candidate_id": "candidate-2",
            "candidates": [
                {"candidate_id": "candidate-1", "status": "draft"},
                {"candidate_id": "candidate-2", "status": "draft"},
            ],
        }

        self.assertEqual(_selected_candidate(task)["candidate_id"], "candidate-2")
        self.assertIsNone(_selected_candidate({**task, "selected_candidate_id": "missing"}))
        self.assertTrue(
            _preview_matches_task(
                {"task_id": "task-1", "snapshot_id": "snapshot-1", "format": "md"},
                task,
                "markdown",
            )
        )
        self.assertFalse(
            _preview_matches_task(
                {"task_id": "task-1", "snapshot_id": "stale", "format": "md"},
                task,
                "markdown",
            )
        )


if __name__ == "__main__":
    unittest.main()
